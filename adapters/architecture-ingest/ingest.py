#!/usr/bin/env python3
"""Untrusted architecture document ingestion. Never executes document content."""
import argparse
import hashlib
import json
from pathlib import Path

SUPPORTED = {'.pdf', '.docx', '.json', '.md', '.txt', '.yaml', '.yml'}
MAX_BYTES = 20 * 1024 * 1024
MAX_CHARS = 2_000_000

def parse(path):
    p = Path(path)
    if p.suffix.lower() not in SUPPORTED:
        raise ValueError('unsupported format; legacy .doc requires a separate isolated converter')
    if not p.is_file() or p.stat().st_size > MAX_BYTES:
        raise ValueError('missing or oversized input')
    data = p.read_bytes()
    suffix = p.suffix.lower()
    sections = []
    def add(loc, text):
        if text and text.strip():
            sections.append({'location': loc, 'text': text[:MAX_CHARS]})
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(str(p), strict=True)
        if reader.is_encrypted:
            raise ValueError('encrypted PDF not supported')
        if len(reader.pages) > 500:
            raise ValueError('too many PDF pages')
        for n, page in enumerate(reader.pages, 1):
            add(f'page:{n}', page.extract_text() or '')
    elif suffix == '.docx':
        from docx import Document
        d = Document(str(p))
        for n, para in enumerate(d.paragraphs, 1):
            add(f'paragraph:{n}', para.text)
        for n, table in enumerate(d.tables, 1):
            for row, cells in enumerate(table.rows, 1):
                add(f'table:{n}:row:{row}', ' | '.join(c.text for c in cells.cells))
    elif suffix == '.json':
        value = json.loads(data.decode('utf-8-sig'))
        add('json:$', json.dumps(value, ensure_ascii=False, sort_keys=True))
    elif suffix in ('.yaml', '.yml'):
        import yaml
        value = yaml.safe_load(data.decode('utf-8-sig'))
        add('yaml:$', json.dumps(value, ensure_ascii=False, sort_keys=True))
    else:
        for n, line in enumerate(data.decode('utf-8-sig').splitlines(), 1):
            add(f'line:{n}', line)
    total = sum(len(s['text']) for s in sections)
    if total > MAX_CHARS:
        raise ValueError('extracted content exceeds character limit')
    return {'schema': 'gauntlet-architecture-source-v0.1', 'source_name': p.name, 'sha256': hashlib.sha256(data).hexdigest(), 'format': suffix, 'sections': sections, 'warnings': ['No readable text extracted; inspect scanned pages or images manually'] if not sections else [], 'assurance': 'UNVERIFIED_SOURCE'}

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('--output')
    args = parser.parse_args()
    result = json.dumps(parse(args.source), ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(result, encoding='utf-8')
    else:
        print(result)

if __name__ == '__main__':
    main()
